from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.title
    if title_box is None:
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(11.5), Inches(0.7))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 51, 88)

    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(11.5), Inches(0.3))
        tf2 = box.text_frame
        tf2.text = subtitle
        tf2.paragraphs[0].alignment = PP_ALIGN.LEFT
        tf2.paragraphs[0].runs[0].font.size = Pt(10)
        tf2.paragraphs[0].runs[0].font.color.rgb = RGBColor(99, 114, 137)


def add_bullets(slide, bullets, left=0.8, top=1.4, width=10.5, height=5.8, font_size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.bullet = True
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        for run in p.runs:
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(44, 62, 80)


slides = [
    {
        'title': '扬州中国大运河博物馆非遗展厅展览叙事研究',
        'subtitle': '课程汇报 · 研究生阶段',
        'bullets': [
            '汇报人：XXX',
            '专业：XXXX',
            '指导老师：XXX',
            '时间：2026年10月',
        ],
    },
    {
        'title': '目录',
        'bullets': [
            '一、选题依据',
            '二、研究背景与文献综述',
            '三、理论框架与研究方法',
            '四、案例分析：扬州中国大运河博物馆非遗展厅',
            '五、存在问题',
            '六、优化路径',
            '七、结语与展望',
            '八、参考文献',
        ],
    },
    {
        'title': '研究背景：为什么做这个课题',
        'bullets': [
            '非遗作为文化资源，逐渐进入博物馆公共文化场域。',
            '传统“器物陈列”已不能满足当代观众对文化理解的需求。',
            '博物馆呈现方式正在从“展示物”转向“讲述故事”。',
            '运河文化与非遗传承结合，突显地方文化记忆的重要性。',
        ],
    },
    {
        'title': '研究对象与切入点',
        'bullets': [
            '研究对象：扬州中国大运河博物馆非遗展厅。',
            '切入点：展厅叙事结构、空间布局、观众体验、文化传播。',
            '研究重点：从叙事视角分析非遗展厅如何讲述地方文化。',
            '目标：探讨展示如何从“静态陈列”转向“文化理解”与“情感认同”。',
        ],
    },
    {
        'title': '选题依据一：政策背景',
        'bullets': [
            '国家高度重视非遗保护与传承。',
            '文化遗产保护与文旅融合成为重要发展方向。',
            '博物馆承担文化教育、传播与传承功能。',
            '非遗进入博物馆成为文化展示的重要趋势。',
        ],
    },
    {
        'title': '选题依据二：地方文化资源优势',
        'bullets': [
            '扬州位于运河文化带核心区域，文化资源丰富。',
            '运河文化具有历史连续性与地方特色。',
            '非遗与地方生活、民俗、技艺相互关联。',
            '非遗展厅是地方文化记忆的重要承载空间。',
        ],
    },
    {
        'title': '选题依据三：博物馆功能转型',
        'bullets': [
            '从“收藏、保存”转向“传播、教育、体验”。',
            '观众逐渐重视观看过程中的情感和认知体验。',
            '展厅不再只是展示物，而是叙事场域。',
            '叙事设计有利于提升非遗传播效果。',
        ],
    },
    {
        'title': '研究意义一：理论意义',
        'bullets': [
            '丰富博物馆叙事理论在非遗展陈中的应用。',
            '扩展地方文化传播与文化记忆研究。',
            '促进“空间叙事”“文化传播”“非遗传承”研究融合。',
            '为运河文化研究提供新的观察视角。',
        ],
    },
    {
        'title': '研究意义二：实践意义',
        'bullets': [
            '提高非遗展厅展示效果和传播力。',
            '增强观众参与感和文化理解力。',
            '为地方博物馆展陈设计提供参考。',
            '推动非遗从“静态收藏”走向“活态传承”。',
        ],
    },
    {
        'title': '国内研究现状：非遗与博物馆',
        'bullets': [
            '研究多集中在非遗保护、文化政策和传播策略。',
            '博物馆中的非遗展示研究逐渐增多。',
            '但对“叙事性”“空间体验”关注仍显不足。',
            '研究通常偏重陈列手段，忽视文化叙事结构。',
        ],
    },
    {
        'title': '国内研究现状：博物馆叙事研究',
        'bullets': [
            '国内学者已关注展陈设计、空间布局和观众体验。',
            '相关研究强调叙事结构和情感传播。',
            '但对地方性非遗展厅的叙事系统性研究较少。',
            '缺乏与运河文化、地方生活结合的分析框架。',
        ],
    },
    {
        'title': '国外研究现状：Museology与Narrative',
        'bullets': [
            '国外博物馆学重视叙事化设计与观众参与。',
            '注重“故事性空间”“体验式展示”“记忆构建”。',
            '对非遗、文化遗产和公共教育结合的研究较深。',
            '可为中国地方博物馆展陈创新提供启发。',
        ],
    },
    {
        'title': '文献综述总结',
        'bullets': [
            '现有研究有：政策研究、展示研究、传播研究。',
            '但缺少：地方性、非遗、运河文化、叙事逻辑的综合研究。',
            '现状问题：展陈重“结果”，轻“过程”；重“物”，轻“人”。',
            '因此，本文在叙事视角下的研究是有必要的。',
        ],
    },
    {
        'title': '研究问题与研究内容',
        'bullets': [
            '研究问题：展厅叙事结构如何呈现？',
            '空间布局和展示手段如何影响观众阅读体验？',
            '该展厅在非遗传播中存在怎样的问题？',
            '如何优化展陈叙事设计？',
        ],
    },
    {
        'title': '理论框架一：博物馆叙事理论',
        'bullets': [
            '展厅可视为“叙事空间”而非单纯“物品陈列室”。',
            '展陈中的时间顺序、空间安排和信息呈现形成叙事结构。',
            '观众在参观过程中进行意义重构和情感体验。',
            '从叙事理念看，展陈是文化解释的过程。',
        ],
    },
    {
        'title': '理论框架二：空间叙事理论',
        'bullets': [
            '空间是叙事的重要媒介。',
            '展厅的动线、节点、光线和标识共同塑造叙事感。',
            '观众通过路径体验感受历史和文化的连续性。',
            '空间设计决定了文化信息的理解顺序和情感强度。',
        ],
    },
    {
        'title': '理论框架三：文化传播理论',
        'bullets': [
            '非遗传播不仅是信息传递，还包括文化认同与情感共鸣。',
            '展陈要让观众理解“非遗为什么重要”。',
            '传播效果依赖展示内容、媒介形式和参观体验。',
            '展厅是文化传播与公共教育的重要载体。',
        ],
    },
    {
        'title': '理论框架四：非遗传承理论',
        'bullets': [
            '非遗强调“活态传承”和“过程性”。',
            '单纯展示器物不足以呈现文化生命。',
            '技艺过程、传承人故事、生活语境更能体现非遗价值。',
            '展陈应突出“看得见的文化”和“理解中的文化”。',
        ],
    },
    {
        'title': '研究方法一：文献分析法',
        'bullets': [
            '对博物馆叙事、非遗展陈、运河文化研究文献进行梳理。',
            '形成研究框架和讨论逻辑。',
            '为案例分析提供理论支撑。',
            '有助于识别研究空白与切入点。',
        ],
    },
    {
        'title': '研究方法二：实地考察法',
        'bullets': [
            '走访扬州中国大运河博物馆非遗展厅。',
            '观察展陈布局、空间动线、展示材料和媒介方式。',
            '记录展示的结构和重点内容。',
            '为问题诊断提供真实材料基础。',
        ],
    },
    {
        'title': '研究方法三：访谈法与体验评估法',
        'bullets': [
            '访谈讲解员、馆内工作人员及参观游客。',
            '了解展陈的传播效果和观众认知差异。',
            '分析展厅是否真正实现文化传达与情感连接。',
            '将“观众感受”纳入评价体系。',
        ],
    },
    {
        'title': '研究方法框架图',
        'bullets': [
            '文献分析：构建叙事理论框架与研究空白。',
            '实地考察：观察展陈空间、动线和信息呈现。',
            '访谈/体验：了解游客认知差异和传播效果。',
            '综合判断：识别问题并提出优化路径。',
        ],
    },
    {
        'title': '案例概况：扬州中国大运河博物馆',
        'bullets': [
            '扬州是运河文化的重要发源与延续地。',
            '博物馆为地方历史文化传播提供平台。',
            '运河文化与地方文旅发展相结合，公共文化价值突出。',
            '非遗展厅体现了地方文化记忆与现代传播需求。',
        ],
    },
    {
        'title': '非遗展厅概况',
        'bullets': [
            '展厅主题围绕运河文化与非遗传承展开。',
            '展品形式包括器物、工艺、技艺、民俗等。',
            '展厅内容呈现具有地域性和文化性。',
            '展陈既有收藏属性，也具传播教育属性。',
        ],
    },
    {
        'title': '展厅叙事结构分析一：整体逻辑',
        'bullets': [
            '展厅整体叙事存在“从历史到文化、从技艺到生活”的逻辑。',
            '叙事设计决定了观众理解文化的路径。',
            '若逻辑不清，容易出现“看完但不理解”的现象。',
            '叙事主线需要明确“时间—空间—经验”关系。',
        ],
    },
    {
        'title': '展厅叙事结构分析二：空间动线',
        'bullets': [
            '展厅动线决定观众进入文化世界的方式。',
            '节点设置影响观众停留时间和注意力。',
            '空间的层次感决定叙事推进是否顺畅。',
            '优秀展陈应使观众“自然进入故事”。',
        ],
    },
    {
        'title': '展厅叙事结构分析三：视觉表达',
        'bullets': [
            '展品说明、图片、影像和互动屏是重要叙事载体。',
            '文字说明是否简洁、准确、易懂影响理解深度。',
            '照片和影像提供情感体验，强化故事感。',
            '视觉信息需要形成统一叙事系统。',
        ],
    },
    {
        'title': '展厅叙事结构分析四：观众体验',
        'bullets': [
            '观众的停留时间、路径选择和关注点可反映叙事效果。',
            '观众容易被“器物外观”吸引，但难以理解文化背景。',
            '叙事设计关键在于让观众理解“为什么它重要”。',
            '体验感是提升传播效果的重要因素。',
        ],
    },
    {
        'title': '存在问题一：叙事逻辑不够清晰',
        'bullets': [
            '展陈往往重物轻故事。',
            '非遗的文化背景信息不足。',
            '观众难以建立“时间—空间—技艺—生活”的完整认知。',
            '导致展厅存在“可见、不可解”的问题。',
        ],
    },
    {
        'title': '存在问题二：空间体验不足',
        'bullets': [
            '动线过于平铺，故事推进感不强。',
            '节点设计缺乏高潮和转换。',
            '观众容易形成“走马观花式参观”。',
            '展厅在情感层面未充分调动观众参与。',
        ],
    },
    {
        'title': '存在问题三：互动与传承表达不足',
        'bullets': [
            '非遗的“活态性”没有被充分体现。',
            '缺少技艺演示、现场体验或参与环节。',
            '展览信息较为静态，传播效率有限。',
            '观众难以感受非遗的生命力和传承价值。',
        ],
    },
    {
        'title': '存在问题四：地方文化认同弱化',
        'bullets': [
            '地方文化符号不够突出。',
            '运河文化的历史脉络与生活经验未充分展现。',
            '可能出现“非遗展示”而非“地方文化传播”的情况。',
            '观众难以形成强烈的文化认同和情感记忆。',
        ],
    },
    {
        'title': '优化路径一：重构叙事主线',
        'bullets': [
            '以“地方历史—技艺传承—生活实践—文化认同”为主线。',
            '把散乱陈列改为具有情节递进的叙事结构。',
            '借助时间线、人物线和空间线增强可读性。',
            '提升展厅的解释性和故事性。',
        ],
    },
    {
        'title': '优化路径二：强化空间叙事设计',
        'bullets': [
            '明确入口、主展区、高潮区和尾声区。',
            '通过路径和节点设计增强文化体验。',
            '提升视觉层次、空间转换和情绪推进。',
            '让观众在空间中自然“理解故事”。',
        ],
    },
    {
        'title': '优化路径三：增强互动性和体验感',
        'bullets': [
            '增加技艺演示、触摸互动、投影展示等形式。',
            '让观众参与非遗知识和工艺过程的理解。',
            '建立“观看—体验—理解”的传播链条。',
            '让非遗从“陈列物”变为“现场文化”。',
        ],
    },
    {
        'title': '优化路径四：提升地方文化表达',
        'bullets': [
            '强化运河文化、扬州生活经验和地方记忆的链接。',
            '用更具体的场景和故事来呈现非遗价值。',
            '增强地方特色与文化认同。',
            '让展厅更像“地方文化的语言”。',
        ],
    },
    {
        'title': '创新点与研究价值',
        'bullets': [
            '从叙事角度理解非遗展厅的传播逻辑。',
            '把“器物陈列”转化为“文化表达与体验空间”。',
            '提供适用于地方博物馆的非遗叙事优化思路。',
            '为展示设计提升提供实践参考。',
        ],
    },
    {
        'title': '结语',
        'bullets': [
            '非遗展厅不仅是展示空间，更是文化传承和认同形成的场域。',
            '叙事设计决定了文化传播是否有效。',
            '未来的展陈应兼顾真实性、体验性和传播性。',
            '让观众在观看中理解文化，在参与中感受文化。',
        ],
    },
    {
        'title': '参考文献',
        'bullets': [
            '文化和旅游部：非物质文化遗产法及相关政策文件。',
            '博物馆学、博物馆叙事与空间设计研究相关论文。',
            '非遗传播、地方文化记忆和运河文化相关文献。',
            '请根据最终论文格式要求统一列出正文引用。',
        ],
    },
    {
        'title': '致谢与提问',
        'bullets': [
            '感谢老师和同学的指导与支持。',
            '本研究从叙事视角分析非遗展厅的展示逻辑与传播问题。',
            '欢迎老师和同学提出宝贵意见。',
            'Q&A',
        ],
    },
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

for slide_data in slides:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # background
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(250, 250, 250)

    # top accent bar
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(30, 90, 150)
    bar.line.fill.background()

    # section number
    if slide_data['title'] != '扬州中国大运河博物馆非遗展厅展览叙事研究':
        num = slide.shapes.add_textbox(Inches(11.2), Inches(0.32), Inches(1.5), Inches(0.4))
        tf = num.text_frame
        tf.text = f"{slides.index(slide_data) + 1}"
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        tf.paragraphs[0].runs[0].font.size = Pt(12)
        tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(80, 100, 130)

    add_title(slide, slide_data['title'], slide_data.get('subtitle'))
    add_bullets(slide, slide_data.get('bullets', []), left=0.9, top=1.5, width=11.3, height=4.8)

# set title slide specific styling with larger title and big subtitle
for idx, slide in enumerate(prs.slides):
    if idx == 0:
        title_box = slide.shapes.title
        if title_box is None:
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.6), Inches(1.2))
        tf = title_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = '扬州中国大运河博物馆非遗展厅展览叙事研究'
        p.alignment = PP_ALIGN.CENTER
        run = p.runs[0]
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 51, 88)

        sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.1), Inches(11.6), Inches(0.5))
        stf = sub.text_frame
        stf.text = '课程汇报 · 研究生阶段'
        stf.paragraphs[0].alignment = PP_ALIGN.CENTER
        stf.paragraphs[0].runs[0].font.size = Pt(16)
        stf.paragraphs[0].runs[0].font.color.rgb = RGBColor(85, 100, 120)

        info = slide.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(11.6), Inches(1.0))
        itf = info.text_frame
        itf.text = '汇报人：XXX\n专业：XXXX\n指导老师：XXX\n时间：2026年10月'
        itf.paragraphs[0].alignment = PP_ALIGN.CENTER
        for paragraph in itf.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(18)
                run.font.color.rgb = RGBColor(44, 62, 80)

        # add decorative line
        line = slide.shapes.add_shape(1, Inches(3.4), Inches(5.8), Inches(6.5), Inches(0.05))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(188, 200, 214)
        line.line.fill.background()

out_dir = 'exports'
import os
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, '扬州中国大运河博物馆非遗展厅展览叙事研究_40页_汇报版.pptx')
prs.save(out_path)
print(f'PPT saved to: {out_path}')
