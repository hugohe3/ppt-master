#!/usr/bin/env python3
"""Diagnose hidden-slide visibility unsupported by template preview export.

Usage:
    python3 -m unittest tests.test_template_hidden_slide_diagnostic

Dependencies:
    python-pptx and the script runtime dependencies.
"""

from __future__ import annotations

import contextlib
import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches

import pptx_template_import
from pptx_to_svg.converter import ConvertOptions, convert_pptx_to_svg


class TemplateHiddenSlideDiagnosticTests(unittest.TestCase):
    def test_template_warns_without_dropping_slide_or_warning_about_roundtrip(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.element.set("show", "0")
            textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            textbox.text = "HIDDEN_SLIDE_MARKER"
            source = root / "source.pptx"
            presentation.save(source)
            warnings = io.StringIO()
            imported = root / "template"
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(warnings), patch.object(
                sys, "argv", ["pptx_template_import.py", str(source), "-o", str(imported)],
            ):
                self.assertEqual(pptx_template_import.main(), 0)
            self.assertIn("source slide 1 (ppt/slides/slide1.xml)", warnings.getvalue())
            self.assertIn("template preview export does not preserve", warnings.getvalue())
            self.assertIn("HIDDEN_SLIDE_MARKER", (imported / "svg/slide_01.svg").read_text(encoding="utf-8"))
            self.assertEqual((imported / "sources/source.pptx").read_bytes(), source.read_bytes())
            native_warnings = io.StringIO()
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(native_warnings):
                converted = convert_pptx_to_svg(
                    source, root / "roundtrip", ConvertOptions(roundtrip=True),
                )
            self.assertEqual(len(converted.slides), 1)
            self.assertNotIn("template preview export", native_warnings.getvalue())
