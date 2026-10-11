#!/usr/bin/env python3
"""PPT Master - DOCX content preservation regressions.

Usage: python3 -m unittest tests.test_source_to_md_content_loss_docx
Dependencies: python-docx, mammoth
"""

import base64
import io
import json
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SCRIPTS_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS_DIR))
sys.path.insert(0, str(SCRIPTS_DIR / "source_to_md"))
import doc_to_md  # noqa: E402


class NotebookContentLossTests(unittest.TestCase):
    def test_attachments_do_not_overwrite_named_outputs_or_each_other(self) -> None:
        try:
            import nbconvert  # noqa: F401
            import nbformat
        except ImportError:
            self.skipTest("optional nbconvert/nbformat unavailable")
        from PIL import Image

        for filename in ("attachment_001", "attachment_001.png"):
            with self.subTest(filename=filename), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                expected = []
                encoded = []
                for color in ("red", "blue", "green"):
                    buffer = io.BytesIO()
                    Image.new("RGB", (8, 8), color).save(buffer, format="PNG")
                    expected.append(buffer.getvalue())
                    encoded.append(base64.b64encode(buffer.getvalue()).decode("ascii"))
                notebook = nbformat.v4.new_notebook(cells=[
                    nbformat.v4.new_code_cell("# Named output", outputs=[
                        nbformat.v4.new_output(
                            "display_data", data={"image/png": encoded[0]},
                            metadata={"filename": filename},
                        ),
                    ]),
                    *[
                        nbformat.v4.new_markdown_cell(
                            "![attachment](attachment:picture.png)",
                            attachments={"picture.png": {"image/png": data}},
                        )
                        for data in encoded[1:]
                    ],
                ])
                source = root / "book.ipynb"
                nbformat.write(notebook, source)
                with redirect_stdout(io.StringIO()):
                    markdown = doc_to_md.convert_to_markdown(str(source))
                refs = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", markdown)
                self.assertEqual(len(refs), 3)
                self.assertEqual(len(set(refs)), 3)
                for ref, raw in zip(refs, expected):
                    self.assertEqual(Path(ref).parent, Path("book_files"))
                    self.assertEqual((root / ref).read_bytes(), raw)
                manifest = json.loads((root / "book_files/image_manifest.json").read_text(encoding="utf-8"))
                self.assertEqual({row["filename"] for row in manifest}, {Path(ref).name for ref in refs})
                self.assertTrue(all(len(row["occurrences"]) == 1 for row in manifest))

    def test_batch_images_and_attachments_remain_independent_after_import(self) -> None:
        try:
            import nbconvert  # noqa: F401
            import nbformat
        except ImportError:
            self.skipTest("optional nbconvert/nbformat unavailable")
        from PIL import Image
        from project_management.cli import ProjectManager

        for same_stem in (False, True):
            with self.subTest(same_stem=same_stem), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                sources = []
                expected = []
                for index, color in enumerate(("red", "blue")):
                    source_dir = root / str(index)
                    source_dir.mkdir()
                    source = source_dir / ("book.ipynb" if same_stem else f"book{index}.ipynb")
                    buffer = io.BytesIO()
                    Image.new("RGB", (8, 8), color).save(buffer, format="PNG")
                    raw = buffer.getvalue()
                    encoded = base64.b64encode(raw).decode("ascii")
                    output = nbformat.v4.new_output(
                        "display_data", data={"image/png": encoded},
                        metadata={"filename": "figure"} if same_stem else {},
                    )
                    notebook = nbformat.v4.new_notebook(cells=[
                        nbformat.v4.new_code_cell(f"# Notebook {index}", outputs=[output]),
                        nbformat.v4.new_markdown_cell(
                            "![attachment](attachment:picture.png)",
                            attachments={"picture.png": {"image/png": encoded}},
                        ),
                    ])
                    nbformat.write(notebook, source)
                    sources.append(source)
                    expected.append(raw)
                output_dir = root / "converted"
                result = subprocess.run(
                    [sys.executable, str(SCRIPTS_DIR / "source_to_md/doc_to_md.py"),
                     *map(str, sources), "-o", str(output_dir)],
                    capture_output=True, text=True, encoding="utf-8", check=False,
                )
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                markdown_paths = sorted(output_dir.glob("*.md"))
                self.assertEqual(len(markdown_paths), 2)
                self.assertEqual(list(output_dir.glob("*.png")), [])
                for markdown, raw in zip(markdown_paths, expected):
                    refs = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", markdown.read_text(encoding="utf-8"))
                    self.assertEqual(len(refs), 2)
                    for ref in refs:
                        self.assertEqual(Path(ref).parent, Path(f"{markdown.stem}_files"))
                        self.assertEqual((output_dir / ref).read_bytes(), raw)
                    profile = json.loads(markdown.with_suffix(".conversion_profile.json").read_text())
                    self.assertEqual(profile["outputs"]["image_count"], 2)
                    manifest = json.loads((output_dir / profile["outputs"]["image_manifest"]).read_text())
                    self.assertEqual(len(manifest), 2)
                manager = ProjectManager(base_dir=root / "projects")
                with redirect_stdout(io.StringIO()):
                    project = Path(manager.init_project("notebooks", quick_generate=True))
                    summary = manager.import_sources(str(project), list(map(str, markdown_paths)), copy=True)
                self.assertEqual(summary["skipped"], [])
                self.assertEqual(len(summary["assets"]), 2)
                for path, raw in zip(summary["markdown"], expected):
                    markdown = Path(path)
                    refs = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", markdown.read_text(encoding="utf-8"))
                    for ref in refs:
                        self.assertEqual((markdown.parent / ref).read_bytes(), raw)
                manifest = json.loads((project / "images/image_manifest.json").read_text())
                self.assertEqual({row["source_namespace"] for row in manifest},
                                 {path.stem for path in markdown_paths})
                self.assertEqual({(project / "images" / row["filename"]).read_bytes() for row in manifest},
                                 set(expected))


class DocxContentLossTests(unittest.TestCase):
    def _convert(self, document, patch_archive=None):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.docx"
            if patch_archive:
                buf = io.BytesIO()
                document.save(buf)
                patch_archive(buf, source)
            else:
                document.save(source)
            with redirect_stdout(io.StringIO()):
                markdown = doc_to_md.convert_to_markdown(str(source))
            profile = json.loads(source.with_suffix(".conversion_profile.json").read_text(encoding="utf-8"))
            return markdown, profile

    def test_f01_nested_table_body_is_preserved(self) -> None:
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Outer body"
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "NESTED_SENTINEL_3000000"
        md, _ = self._convert(doc)
        self.assertIn("Outer body", md)
        self.assertIn("NESTED_SENTINEL_3000000", md)

    def test_f02_content_control_display_text_is_preserved(self) -> None:
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "Before"
        cell._tc.append(parse_xml('<w:sdt ' + nsdecls('w') + '><w:sdtContent><w:p><w:r>'
                                 '<w:t>CONTROL_BODY</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
        md, profile = self._convert(doc)
        self.assertIn("CONTROL_BODY", md)
        self.assertTrue(any("content control" in w.lower() for w in profile["warnings"]))

    def test_f03_table_endnote_reference_and_body_are_preserved(self) -> None:
        doc = Document()
        p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        p.add_run("Reference")
        p._p.append(parse_xml('<w:r ' + nsdecls('w') + '><w:endnoteReference w:id="1"/></w:r>'))

        def patch_archive(buf, source):
            with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(source, "w") as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename == "word/_rels/document.xml.rels":
                        root = ET.fromstring(data)
                        ET.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/relationships}'
                                      'Relationship', Id="rIdEndnotes", Type=RT.ENDNOTES, Target="endnotes.xml")
                        data = ET.tostring(root)
                    if item.filename == "[Content_Types].xml":
                        root = ET.fromstring(data)
                        ET.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/content-types}'
                                      'Override', PartName="/word/endnotes.xml", ContentType=
                                      "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml")
                        data = ET.tostring(root)
                    zout.writestr(item, data)
                zout.writestr("word/endnotes.xml", '<w:endnotes ' + nsdecls('w') + '><w:endnote w:id="1">'
                              '<w:p><w:r><w:t>ENDNOTE_BODY</w:t></w:r></w:p></w:endnote></w:endnotes>')

        md, _ = self._convert(doc, patch_archive)
        self.assertIn("[^endnote-1]", md)
        self.assertIn("[^endnote-1]: ENDNOTE_BODY", md)

    def test_f09_table_hyperlink_target_is_preserved(self) -> None:
        doc = Document()
        p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        rid = p.part.relate_to("https://example.invalid/evidence", RT.HYPERLINK, is_external=True)
        p._p.append(parse_xml('<w:hyperlink ' + nsdecls('w', 'r') + ' r:id="' + rid + '">'
                             '<w:r><w:t>Evidence</w:t></w:r></w:hyperlink>'))
        md, _ = self._convert(doc)
        self.assertIn("| [Evidence](https://example.invalid/evidence) |", md)
        self.assertIn("| --- |", md)

    def _math_document(self, xml):
        doc = Document()
        doc.add_paragraph()._p.append(parse_xml('<m:oMath ' + nsdecls('m') + '>' + xml + '</m:oMath>'))
        return doc

    def test_f23_delimiter_operands_keep_separator(self) -> None:
        for properties, separator in [('<m:dPr><m:sepChr m:val="|"/></m:dPr>', '|'), ('', '|'),
                                      ('<m:dPr><m:sepChr m:val=";"/></m:dPr>', ';')]:
            with self.subTest(separator=separator, properties=properties):
                md, _ = self._convert(self._math_document('<m:d>' + properties +
                    '<m:e><m:r><m:t>x</m:t></m:r></m:e><m:e><m:r><m:t>y</m:t></m:r></m:e></m:d>'))
                self.assertIn('x' + separator + 'y', md)

    def test_f24_group_character_and_position_are_preserved(self) -> None:
        for char, pos, expected in [('⏞', 'top', r'\overbrace{x+y}'),
                                    ('⏟', 'bot', r'\underbrace{x+y}')]:
            with self.subTest(char=char):
                md, _ = self._convert(self._math_document('<m:groupChr><m:groupChrPr><m:chr m:val="' + char +
                    '"/><m:pos m:val="' + pos + '"/></m:groupChrPr><m:e><m:r><m:t>x+y</m:t></m:r>'
                    '</m:e></m:groupChr>'))
                self.assertIn(expected, md)

    def test_f24_unknown_group_character_is_retained_with_warning(self) -> None:
        md, profile = self._convert(self._math_document('<m:groupChr><m:groupChrPr><m:chr m:val="?"/>'
            '</m:groupChrPr><m:e><m:r><m:t>x</m:t></m:r></m:e></m:groupChr>'))
        self.assertIn('?', md)
        self.assertTrue(any('group character' in w for w in profile['warnings']))
