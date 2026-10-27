#!/usr/bin/env python3
"""Report intentional hidden-shape exclusions without changing source objects.

Usage:
    python3 -m unittest tests.test_pptx_hidden_shape_diagnostics

Dependencies:
    python-pptx and the script runtime dependencies.
"""

from __future__ import annotations

import contextlib
import io
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches

from pptx_to_svg.converter import ConvertOptions, convert_pptx_to_svg


class HiddenShapeDiagnosticTests(unittest.TestCase):
    def test_hidden_shape_and_group_are_excluded_with_diagnostics(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            visible = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            visible.text = "VISIBLE_MARKER"
            hidden = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
            hidden.text = "HIDDEN_SHAPE_MARKER"
            hidden.element.xpath(".//p:cNvPr")[0].set("hidden", "1")
            group = slide.shapes.add_group_shape()
            child = group.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
            child.text = "HIDDEN_GROUP_MARKER"
            group.element.xpath("./p:nvGrpSpPr/p:cNvPr")[0].set("hidden", "true")
            source = root / "source.pptx"
            presentation.save(source)
            before = source.read_bytes()
            for keep_hidden in (False, True):
                destination = root / str(keep_hidden)
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    result = convert_pptx_to_svg(
                        source, destination,
                        ConvertOptions(inheritance_mode="both", keep_hidden=keep_hidden),
                    )
                text = (destination / "svg-flat/slide_01.svg").read_text(encoding="utf-8")
                self.assertIn("VISIBLE_MARKER", text)
                warnings = [d for d in result.diagnostics if d.code == "author-hidden-omitted"]
                if keep_hidden:
                    self.assertIn("HIDDEN_SHAPE_MARKER", text)
                    self.assertIn("HIDDEN_GROUP_MARKER", text)
                    self.assertEqual(warnings, [])
                else:
                    self.assertNotIn("HIDDEN_SHAPE_MARKER", text)
                    self.assertNotIn("HIDDEN_GROUP_MARKER", text)
                    self.assertEqual({d.shape_id for d in warnings}, {str(hidden.shape_id), str(group.shape_id)})
                    self.assertTrue(all(d.part_path == "ppt/slides/slide1.xml" for d in warnings))
            self.assertEqual(source.read_bytes(), before)
