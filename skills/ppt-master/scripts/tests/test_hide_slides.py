#!/usr/bin/env python3
"""Unit tests for hidden-slide marking on finished PPTX files."""

from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from xml.etree import ElementTree as ET


SCRIPTS_DIR = Path(__file__).resolve().parents[1]
if not sys.path or sys.path[0] != str(SCRIPTS_DIR):
    sys.path.insert(0, str(SCRIPTS_DIR))

from hide_slides import (  # noqa: E402
    apply_visibility,
    last_pages,
    main,
    pages_matching_stem,
    parse_page_spec,
    set_slide_visibility,
    slide_is_visible,
)
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402
from pptx_delivery_check import audit_pptx_delivery  # noqa: E402
from pptx_ooxml.ooxml import _parse_slide_refs  # noqa: E402


def _write_deck(path: Path, titles: list[str]) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    for title in titles:
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11), Inches(1.2))
        box.text_frame.text = title
    presentation.save(str(path))
    return path


class ParseTests(unittest.TestCase):
    def test_page_spec_ranges_and_lists(self) -> None:
        self.assertEqual(parse_page_spec("5,6,8-10", 12), [5, 6, 8, 9, 10])
        self.assertEqual(last_pages(6, 2), [5, 6])

    def test_page_spec_rejects_out_of_range(self) -> None:
        with self.assertRaises(ValueError):
            parse_page_spec("9", 4)

    def test_match_reads_svg_stems(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            project = Path(tmp)
            svg = project / "svg_output"
            svg.mkdir()
            for name in ("01_cover.svg", "02_body.svg", "03_appendix_a.svg", "04_appendix_b.svg"):
                (svg / name).write_text("<svg xmlns='http://www.w3.org/2000/svg'/>", encoding="utf-8")
            self.assertEqual(pages_matching_stem(project, "appendix"), [3, 4])


class VisibilityXmlTests(unittest.TestCase):
    def test_set_and_clear_show_on_prefixed_root(self) -> None:
        original = (
            b'<?xml version="1.0" encoding="UTF-8"?>'
            b'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            b"<p:cSld/></p:sld>"
        )
        hidden = set_slide_visibility(original, visible=False)
        self.assertFalse(slide_is_visible(hidden))
        self.assertIn(b'show="0"', hidden)
        shown = set_slide_visibility(hidden, visible=True)
        self.assertTrue(slide_is_visible(shown))
        self.assertNotIn(b"show=", shown)


class PackageTests(unittest.TestCase):
    def test_hide_last_two_sets_slide_show_and_delivery_check(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = _write_deck(
                Path(tmp) / "source.pptx",
                ["Cover", "Body", "Appendix A", "Appendix B"],
            )
            dest = Path(tmp) / "hidden.pptx"
            report = apply_visibility(source, dest, [3, 4], visible=False)
            self.assertEqual(report["changed"], [3, 4])
            with zipfile.ZipFile(dest) as package:
                refs = _parse_slide_refs(package)
                vis = [
                    slide_is_visible(package.read(ref.part_name))
                    for ref in refs
                ]
            self.assertEqual(vis, [True, True, False, False])
            delivery = audit_pptx_delivery(dest)
            hidden = delivery["slides"]["hidden"]
            self.assertEqual([item["index"] for item in hidden], [3, 4])

    def test_unhide_restores_default_visibility(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = _write_deck(Path(tmp) / "source.pptx", ["A", "B"])
            hidden = Path(tmp) / "hidden.pptx"
            apply_visibility(source, hidden, [2], visible=False)
            restored = Path(tmp) / "restored.pptx"
            apply_visibility(hidden, restored, [2], visible=True)
            with zipfile.ZipFile(restored) as package:
                refs = _parse_slide_refs(package)
                self.assertTrue(slide_is_visible(package.read(refs[1].part_name)))
            self.assertEqual(audit_pptx_delivery(restored)["slides"]["hidden"], [])


class CliTests(unittest.TestCase):
    def test_cli_match_and_json(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            project = root / "project"
            (project / "svg_output").mkdir(parents=True)
            for name in ("01_cover.svg", "02_body.svg", "03_appendix.svg"):
                (project / "svg_output" / name).write_text(
                    "<svg xmlns='http://www.w3.org/2000/svg'/>", encoding="utf-8"
                )
            source = _write_deck(root / "deck.pptx", ["Cover", "Body", "Appendix"])
            dest = root / "out.pptx"
            stdout = io.StringIO()
            stderr = io.StringIO()
            with redirect_stdout(stdout), redirect_stderr(stderr):
                code = main([
                    str(source),
                    "--project", str(project),
                    "--match", "appendix",
                    "-o", str(dest),
                    "--json",
                ])
            self.assertEqual(code, 0, stderr.getvalue())
            payload = json.loads(stdout.getvalue())
            self.assertEqual(payload["changed"], [3])
            with zipfile.ZipFile(dest) as package:
                refs = {ref.index: ref for ref in _parse_slide_refs(package)}
                self.assertFalse(slide_is_visible(package.read(refs[3].part_name)))

    def test_help_exits_zero(self) -> None:
        stdout = io.StringIO()
        with redirect_stdout(stdout), self.assertRaises(SystemExit) as caught:
            main(["--help"])
        self.assertEqual(caught.exception.code, 0)
        self.assertIn("Hide or unhide", stdout.getvalue())

    def test_missing_pptx_is_nonzero(self) -> None:
        stderr = io.StringIO()
        with redirect_stderr(stderr):
            code = main(["/tmp/missing-hide-slides.pptx", "--pages", "1"])
        self.assertEqual(code, 1)
        self.assertIn("not found", stderr.getvalue())


class ExportThenHideTests(unittest.TestCase):
    def test_hide_after_svg_export(self) -> None:
        import subprocess

        from svg_to_pptx.pptx_package.cli import main as export_main

        page = (
            '<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" '
            'viewBox="0 0 1280 720" data-pptx-page-role="{role}" font-family="Arial">'
            '<rect x="0" y="0" width="1280" height="720" fill="#F8FAFC"/>'
            '<text x="80" y="200" font-family="Arial" font-size="40" fill="#0F172A">'
            "{title}</text></svg>\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            project = Path(tmp) / "project"
            svg = project / "svg_output"
            svg.mkdir(parents=True)
            (svg / "01_cover.svg").write_text(page.format(role="cover", title="Cover"), encoding="utf-8")
            (svg / "02_body.svg").write_text(page.format(role="content", title="Body"), encoding="utf-8")
            (svg / "03_appendix.svg").write_text(
                page.format(role="content", title="Appendix"), encoding="utf-8"
            )
            check = subprocess.run(
                [
                    sys.executable, "-B",
                    str(SCRIPTS_DIR / "svg_quality_checker.py"),
                    str(project),
                    "--quick-generate", "--canonical-authoring",
                    "--stage", "final", "--json",
                ],
                capture_output=True, text=True, encoding="utf-8", errors="replace",
            )
            self.assertEqual(check.returncode, 0, check.stdout + check.stderr)
            exported = Path(tmp) / "exported.pptx"
            stderr = io.StringIO()
            stdout = io.StringIO()
            with redirect_stdout(stdout), redirect_stderr(stderr):
                code = export_main([str(project), "--quick-generate", "-o", str(exported)])
            self.assertEqual(code, 0, stdout.getvalue() + stderr.getvalue())
            hidden = Path(tmp) / "hidden.pptx"
            self.assertEqual(
                main([str(exported), "--last", "1", "-o", str(hidden)]),
                0,
            )
            delivery = audit_pptx_delivery(hidden)
            self.assertEqual([item["index"] for item in delivery["slides"]["hidden"]], [3])
            slide = ET.fromstring(
                zipfile.ZipFile(hidden).read("ppt/slides/slide3.xml")
            )
            self.assertEqual(slide.get("show"), "0")


if __name__ == "__main__":
    unittest.main()
