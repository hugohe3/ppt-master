#!/usr/bin/env python3
"""Regressions for the PPTX design-profile extractor.

Usage:
    cd skills/ppt-master/scripts && python3 -m unittest tests.test_pptx_design_extractor

Dependencies:
    Standard library, python-pptx, and the script runtime dependencies.
"""

from __future__ import annotations

import contextlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1]
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

import pptx_design_extractor as extractor  # noqa: E402


def _build_minimal_deck(path: Path) -> None:
    """Write a small deck with known theme, fill, and text properties."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "测试标题"
    cover.placeholders[1].text = "SUBTITLE"

    content = prs.slides.add_slide(prs.slide_layouts[6])
    shape = content.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF9, 0xC3, 0x7A)
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "内容页"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(0x19, 0x21, 0x2C)

    toc = prs.slides.add_slide(prs.slide_layouts[1])
    toc.shapes.title.text = "目录"
    toc.placeholders[1].text = "Agenda"

    prs.save(str(path))


class DesignProfileTest(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)
        self.pptx_path = self.root / "deck.pptx"
        _build_minimal_deck(self.pptx_path)

    def test_extracts_canvas_theme_and_structure(self) -> None:
        profile = extractor.extract_design_profile(self.pptx_path, self.root / "out")

        self.assertEqual(profile["slideSize"]["width_px"], 1280)
        self.assertEqual(profile["slideSize"]["height_px"], 720)
        # The default python-pptx template ships the Office color scheme.
        self.assertIn("accent1", profile["theme"]["colors"])
        self.assertEqual(profile["slideComposition"]["totalSlides"], 3)
        self.assertTrue(profile["structure"]["masters"])
        self.assertTrue(profile["structure"]["layouts"])

    def test_aggregates_fill_color_and_typography(self) -> None:
        profile = extractor.extract_design_profile(self.pptx_path, self.root / "out")

        colors = {item["color"] for item in profile["colorPalette"]["usageFrequency"]}
        self.assertIn("#F9C37A", colors)

        text_colors = {item["color"] for item in profile["colorPalette"]["textColors"]}
        self.assertIn("#19212C", text_colors)

        fonts = {item["font"] for item in profile["typography"]["usageFrequency"]}
        self.assertIn("Arial", fonts)

        sizes = {item["size_pt"] for item in profile["typography"]["fontSizeHierarchy"]}
        self.assertIn(24.0, sizes)

    def test_classifies_cover_toc_content(self) -> None:
        profile = extractor.extract_design_profile(self.pptx_path, self.root / "out")

        by_index = {
            slide["index"]: slide["pageType"]
            for slide in profile["slideComposition"]["slides"]
        }
        self.assertEqual(by_index[1], "cover")
        self.assertEqual(by_index[2], "content")
        self.assertEqual(by_index[3], "toc")

    def test_cli_writes_profile_json(self) -> None:
        out_dir = self.root / "cli_out"
        with contextlib.redirect_stdout(io.StringIO()):
            exit_code = extractor.main([str(self.pptx_path), "-o", str(out_dir)])
        self.assertEqual(exit_code, 0)
        profile_path = out_dir / "design_profile.json"
        self.assertTrue(profile_path.is_file())
        payload = json.loads(profile_path.read_text(encoding="utf-8"))
        self.assertEqual(payload["source_name"], "deck.pptx")
        # Screenshot generation is opt-in; the base run records no backend.
        self.assertNotIn("screenshots", payload)


if __name__ == "__main__":
    unittest.main()
