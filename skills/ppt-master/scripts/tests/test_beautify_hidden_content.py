#!/usr/bin/env python3
"""Exclude author-hidden objects from both Beautify content-source paths.

Usage:
    python3 -m unittest tests.test_beautify_hidden_content

Dependencies:
    python-pptx, Pillow, and the script runtime dependencies.
"""

from __future__ import annotations

import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parents[1]
for directory in (SCRIPTS, SCRIPTS / "source_to_md"):
    if str(directory) not in sys.path:
        sys.path.insert(0, str(directory))

import beautify_inventory  # noqa: E402
import ppt_to_md  # noqa: E402
import pptx_intake  # noqa: E402

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
}


class BeautifyHiddenContentTests(unittest.TestCase):
    def test_hidden_subtrees_are_absent_from_markdown_inventory_and_images(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])

            def textbox(shapes, text):
                shape = shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
                shape.text = text
                return shape

            def hide(shape, value="1"):
                shape.element.xpath("./*/p:cNvPr")[0].set("hidden", value)

            visible = textbox(slide.shapes, "VISIBLE_FALSE_MARKER")
            hide(visible, "false")
            hide(textbox(slide.shapes, "SECRET_TEXT_MARKER"))
            group = slide.shapes.add_group_shape()
            textbox(group.shapes, "SECRET_GROUP_MARKER")
            hide(group, "true")
            table = slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(4), Inches(1))
            table.table.cell(0, 0).text = "SECRET_TABLE_MARKER"
            hide(table)
            data = ChartData()
            data.categories = ["SECRET_CATEGORY_MARKER"]
            data.add_series("SECRET_SERIES_MARKER", [271])
            chart = slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED,
                Inches(1), Inches(1), Inches(8), Inches(4), data,
            )
            chart.chart.has_title = True
            chart.chart.chart_title.text_frame.text = "SECRET_TITLE_MARKER"
            hide(chart)
            image = root / "picture.png"
            Image.new("RGB", (30, 30), "red").save(image)
            slide.shapes.add_picture(str(image), Inches(1), Inches(5))
            hidden_picture = slide.shapes.add_picture(str(image), Inches(2), Inches(5))
            hidden_picture.element.xpath("./*/p:cNvPr")[0].set("descr", "SECRET_IMAGE_MARKER")
            hide(hidden_picture, "on")
            source = root / "source.pptx"
            presentation.save(source)

            with zipfile.ZipFile(source) as archive:
                entries = {name: archive.read(name) for name in archive.namelist()}
            slide_xml = ET.fromstring(entries["ppt/slides/slide1.xml"])
            group_xml = slide_xml.find(".//p:grpSp", NS)
            group_xml.append(ET.fromstring(
                f'<p:graphicFrame xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}" '
                f'xmlns:r="{NS["r"]}" xmlns:dgm="{NS["dgm"]}">'
                '<p:nvGraphicFramePr><p:cNvPr id="90" name="Hidden diagram"/>'
                '<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
                '<p:xfrm><a:off x="0" y="0"/><a:ext cx="1000000" cy="1000000"/></p:xfrm>'
                f'<a:graphic><a:graphicData uri="{NS["dgm"]}">'
                '<dgm:relIds r:dm="rIdHiddenDiagram"/></a:graphicData></a:graphic>'
                '</p:graphicFrame>'
            ))
            entries["ppt/slides/slide1.xml"] = ET.tostring(slide_xml)
            rel_name = "ppt/slides/_rels/slide1.xml.rels"
            rels = ET.fromstring(entries[rel_name])
            rel_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
            ET.SubElement(rels, f"{{{rel_namespace}}}Relationship", {
                "Id": "rIdHiddenDiagram", "Type": NS["r"] + "/diagramData",
                "Target": "../diagrams/data1.xml",
            })
            entries[rel_name] = ET.tostring(rels)
            entries["ppt/diagrams/data1.xml"] = (
                f'<dgm:dataModel xmlns:dgm="{NS["dgm"]}" xmlns:a="{NS["a"]}">'
                '<dgm:ptLst><dgm:pt modelId="1" type="node"><dgm:prSet/>'
                '<dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r>'
                '<a:t>SECRET_DIAGRAM_MARKER</a:t></a:r></a:p></dgm:t>'
                '</dgm:pt></dgm:ptLst><dgm:cxnLst/></dgm:dataModel>'
            ).encode()
            content_types = ET.fromstring(entries["[Content_Types].xml"])
            ET.SubElement(content_types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
                "PartName": "/ppt/diagrams/data1.xml",
                "ContentType": "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
            })
            entries["[Content_Types].xml"] = ET.tostring(content_types)
            with zipfile.ZipFile(source, "w") as archive:
                for name, payload in entries.items():
                    archive.writestr(name, payload)
            original = source.read_bytes()
            warnings = io.StringIO()
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(warnings):
                self.assertEqual(ppt_to_md.main([str(source), "-o", str(root / "source.md")]), 0)
                self.assertEqual(pptx_intake.main([str(source), "-o", str(root / "analysis")]), 0)
                self.assertEqual(beautify_inventory.main([
                    str(root / "analysis/source.slide_library.json"),
                    "--images", str(root / "source_files/image_manifest.json"),
                    "-o", str(root / "inventory.json"),
                ]), 0)
            markdown = (root / "source.md").read_text(encoding="utf-8")
            inventory_text = (root / "inventory.json").read_text(encoding="utf-8")
            for text in (markdown, inventory_text):
                self.assertIn("VISIBLE_FALSE_MARKER", text)
                self.assertNotIn("SECRET_", text)
            inventory = json.loads(inventory_text)
            self.assertEqual(len(inventory["slides"][0]["images"]), 1)
            self.assertEqual(inventory["slides"][0]["tables"], [])
            self.assertEqual(inventory["slides"][0]["charts"], [])
            self.assertEqual(inventory["slides"][0]["diagrams"], [])
            self.assertIn("excluded author-hidden", warnings.getvalue())
            self.assertEqual(source.read_bytes(), original)
