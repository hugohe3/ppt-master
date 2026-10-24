#!/usr/bin/env python3
"""
PPT Master - Imported Chart In-place Data Regression Tests

Exercise public import, checker and round-trip export on synthetic native charts
with source-only formatting and embedded workbook data.

Usage:
    python3 -m unittest tests.test_chart_inplace

Examples:
    Run the module from the scripts directory.

Dependencies:
    python-pptx and the local importer, checker and exporter.
"""

from __future__ import annotations

import copy
import io
import json
import posixpath
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from pptx_to_svg.converter import ConvertOptions, convert_pptx_to_svg
from svg_quality.checker import SVGQualityChecker
from svg_to_pptx.pptx_package.cli import main as export_main


NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "svg": "http://www.w3.org/2000/svg",
}
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


class ChartInplaceTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.workspace = self.root / "workspace"

    def _import(self, *, values=(12, 34), shared_slide: bool = False, named: bool = True) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["First", "Second"]
        data.add_series("Revenue", values)
        frame = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data,
        )
        if shared_slide:
            second = presentation.slides.add_slide(presentation.slide_layouts[6])
            clone = copy.deepcopy(frame._element)
            rel_id = second.part.relate_to(frame.chart.part, NS["r"] + "/chart")
            clone.find(".//c:chart", NS).set(f'{{{NS["r"]}}}id', rel_id)
            second.shapes._spTree.insert_element_before(clone, "p:extLst")
        original = self.root / "original.pptx"
        presentation.save(original)
        with ZipFile(original) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        chart = ET.fromstring(parts["ppt/charts/chart1.xml"])
        series = chart.find(".//c:ser", NS)
        point = ET.fromstring(
            f'<c:dPt xmlns:c="{NS["c"]}" xmlns:a="{NS["a"]}"><c:idx val="0"/>'
            '<c:spPr><a:gradFill><a:gsLst><a:gs pos="0"><a:srgbClr val="112233"/>'
            '</a:gs><a:gs pos="100000"><a:srgbClr val="445566"/></a:gs></a:gsLst>'
            '<a:lin ang="5400000" scaled="1"/></a:gradFill>'
            '<a:ln w="19050"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:ln>'
            '</c:spPr></c:dPt>'
        )
        series.insert(3, point)
        series.insert(3, copy.deepcopy(point.find("c:spPr", NS)))
        legend = ET.fromstring(
            f'<c:legend xmlns:c="{NS["c"]}" xmlns:a="{NS["a"]}"><c:legendPos val="t"/>'
            '<c:legendEntry><c:idx val="0"/><c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr>'
            '<a:defRPr sz="2200"/></a:pPr><a:endParaRPr/></a:p></c:txPr></c:legendEntry>'
            '<c:layout><c:manualLayout><c:x val="0.2"/><c:y val="0.1"/>'
            '<c:w val="0.7"/><c:h val="0.12"/></c:manualLayout></c:layout></c:legend>'
        )
        chart.find("c:chart", NS).append(legend)
        chart.insert(0, ET.Element(f'{{{NS["c"]}}}style', {"val": "10"}))
        extension = ET.SubElement(ET.SubElement(chart, f'{{{NS["c"]}}}extLst'), f'{{{NS["c"]}}}ext',
                                  {"uri": "{synthetic-preserved-extension}"})
        ET.SubElement(extension, "{urn:chart-test}opaque", {"keep": "exact"}).text = "extension data"
        for value in chart.findall(".//c:formatCode", NS):
            value.text = "#,##0.00"
        if not named:
            series.remove(series.find("c:tx", NS))
        # Serialize once before import: these are the exact bytes export must retain.
        parts["ppt/charts/chart1.xml"] = ET.tostring(chart, encoding="utf-8", xml_declaration=True)
        relationships = ET.fromstring(parts["ppt/charts/_rels/chart1.xml.rels"])
        types = ET.fromstring(parts["[Content_Types].xml"])
        for rid, filename, kind, content_type, content in (
            ("rStyle", "style1.xml", "http://schemas.microsoft.com/office/2011/relationships/chartStyle",
             "application/vnd.ms-office.chartstyle+xml",
             '<cs:chartStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle" id="10"/>'),
            ("rColors", "colors1.xml", "http://schemas.microsoft.com/office/2011/relationships/chartColorStyle",
             "application/vnd.ms-office.chartcolorstyle+xml",
             '<cs:colorStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle" id="10"/>'),
            ("rTheme", "../theme/themeOverride1.xml", NS["r"] + "/themeOverride",
             "application/vnd.openxmlformats-officedocument.themeOverride+xml",
             f'<a:themeOverride xmlns:a="{NS["a"]}"/>'),
        ):
            name = posixpath.normpath(posixpath.join("ppt/charts", filename))
            parts[name] = content.encode()
            ET.SubElement(relationships, f"{{{REL_NS}}}Relationship", {"Id": rid, "Type": kind, "Target": filename})
            ET.SubElement(types, f"{{{CT_NS}}}Override", {"PartName": "/" + name, "ContentType": content_type})
        parts["ppt/charts/_rels/chart1.xml.rels"] = ET.tostring(relationships)
        parts["[Content_Types].xml"] = ET.tostring(types)
        self.source = parts
        source = self.root / "source.pptx"
        with ZipFile(source, "w") as archive:
            for name, content in parts.items():
                archive.writestr(name, content)
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            convert_pptx_to_svg(source, self.workspace, ConvertOptions(inheritance_mode="both", roundtrip=True))
        self.path = self.workspace / "authoring-svg-flat" / "slide_01.svg"
        self.svg = ET.parse(self.path).getroot()
        self.marker = next(node for node in self.svg.iter() if node.get("data-pptx-replace-with") == "chart")
        self.metadata = self.marker.find('svg:metadata[@type="application/json"]', NS)
        self.payload = json.loads(self.metadata.text)

    def _save(self) -> None:
        if json.loads(self.metadata.text) != self.payload:
            self.metadata.text = json.dumps(self.payload, ensure_ascii=False)
        self.path.write_bytes(ET.tostring(self.svg))

    def _check(self) -> list:
        self._save()
        checker = SVGQualityChecker()
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            results = checker.check_roundtrip_workspace(str(self.workspace))
        self.assertTrue(results)
        self.assertEqual(results[0]["errors"], [])
        return results[0]["warnings"]

    def _export(self) -> dict[str, bytes]:
        self._save()
        log = io.StringIO()
        target = self.root / "result.pptx"
        with redirect_stdout(log), redirect_stderr(log):
            status = export_main([str(self.workspace), "--roundtrip", "--native-charts-and-tables", "-o", str(target)])
        self.assertEqual(status, 0, log.getvalue())
        self.log = log.getvalue()
        with ZipFile(target) as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    def _live_chart(self, parts: dict[str, bytes], slide: int = 1) -> tuple[str, ET.Element]:
        page = ET.fromstring(parts[f"ppt/slides/slide{slide}.xml"])
        rel_id = page.find(".//c:chart", NS).get(f'{{{NS["r"]}}}id')
        rels = ET.fromstring(parts[f"ppt/slides/_rels/slide{slide}.xml.rels"])
        target = next(rel.get("Target") for rel in rels if rel.get("Id") == rel_id)
        name = posixpath.normpath(posixpath.join("ppt/slides", target))
        return name, ET.fromstring(parts[name])

    def _workbook(self, parts: dict[str, bytes], chart_part: str) -> dict[str, bytes]:
        rels_name = posixpath.join(posixpath.dirname(chart_part), "_rels", posixpath.basename(chart_part) + ".rels")
        rels = ET.fromstring(parts[rels_name])
        target = next(rel.get("Target") for rel in rels if rel.get("Type").endswith("/package"))
        name = posixpath.normpath(posixpath.join(posixpath.dirname(chart_part), target))
        with ZipFile(io.BytesIO(parts[name])) as workbook:
            return {name: workbook.read(name) for name in workbook.namelist()}

    @staticmethod
    def _cells(workbook: dict[str, bytes]) -> dict:
        strings = ET.fromstring(workbook["xl/sharedStrings.xml"])
        values = ["".join(node.itertext()) for node in strings.findall("s:si", NS)]
        cells = {}
        for cell in ET.fromstring(workbook["xl/worksheets/sheet1.xml"]).findall(".//s:c", NS):
            value = cell.find("s:v", NS)
            text = value.text if value is not None else None
            cells[cell.get("r")] = values[int(text)] if cell.get("t") == "s" else text
        return cells

    def test_data_only_preserves_chart_bytes_and_dependency_graph(self) -> None:
        self._import(shared_slide=True)
        self.payload["series"][0].update(name="Income & costs", values=[21, 43])
        self.payload["categories"] = ["新类别", "Second"]
        self.assertEqual(self._check(), [])
        parts = self._export()
        name, chart = self._live_chart(parts)
        # All differences are in existing cache text; namespaces, gradient, legend,
        # c:style and extLst are not reserialized.
        expected = self.source["ppt/charts/chart1.xml"]
        for old, new in ((b">12<", b">21<"), (b">34<", b">43<"), (b">First<", ">新类别<".encode()),
                         (b">Revenue<", b">Income &amp; costs<")):
            expected = expected.replace(old, new)
        self.assertEqual(parts[name], expected)
        self.assertNotEqual(name, "ppt/charts/chart1.xml")
        rels_name = posixpath.join(posixpath.dirname(name), "_rels", posixpath.basename(name) + ".rels")
        for rel in ET.fromstring(parts[rels_name]):
            target = posixpath.normpath(posixpath.join(posixpath.dirname(name), rel.get("Target")))
            self.assertIn(target, parts)
            if rel.get("Id") in {"rStyle", "rColors", "rTheme"}:
                old = next(item.get("Target") for item in ET.fromstring(self.source["ppt/charts/_rels/chart1.xml.rels"])
                           if item.get("Id") == rel.get("Id"))
                self.assertEqual(parts[target], self.source[posixpath.normpath(posixpath.join("ppt/charts", old))])
        workbook = self._workbook(parts, name)
        self.assertEqual(self._cells(workbook), {"B1": "Income & costs", "A2": "新类别", "B2": "21",
                                                "A3": "Second", "B3": "43"})
        self.assertEqual(chart.find(".//c:formatCode", NS).text, "#,##0.00")
        self.assertEqual(self._live_chart(parts, 2)[0], "ppt/charts/chart1.xml")
        self.assertEqual(parts["ppt/charts/chart1.xml"], self.source["ppt/charts/chart1.xml"])
        original_workbook = self._workbook(self.source, "ppt/charts/chart1.xml")
        for part, data in original_workbook.items():
            if part not in {"xl/worksheets/sheet1.xml", "xl/sharedStrings.xml"}:
                self.assertEqual(workbook[part], data)

    def test_series_change_rebuilds_and_warns_with_import_diagnostics(self) -> None:
        self._import()
        self.payload["series"].append({"name": "Added", "values": [7, 8]})
        warning = "\n".join(self._check())
        self.assertIn("native-chart-rebuilt", warning)
        self.assertIn("chart-series-style-normalized", warning)
        self.assertIn("chart-legend-entry-font-size-omitted", warning)
        self.assertIn(self.marker.get("data-pptx-source-ref"), warning)
        parts = self._export()
        name, chart = self._live_chart(parts)
        self.assertNotIn("_data_", name)
        self.assertEqual(len(chart.findall(".//c:ser", NS)), 2)
        self.assertIn("native-chart-rebuilt", self.log)

    def test_non_data_field_change_does_not_patch(self) -> None:
        self._import()
        self.payload["style"]["colors"] = ["#ABCDEF"]
        self.payload["series"][0]["values"] = [21, 43]
        self.assertIn("native-chart-rebuilt", "\n".join(self._check()))
        parts = self._export()
        name, chart = self._live_chart(parts)
        self.assertNotIn("_data_", name)
        self.assertIsNone(chart.find(".//a:gradFill", NS))

    def test_unchanged_chart_is_exact_passthrough(self) -> None:
        self._import()
        # Semantically unchanged authoring must restore the original package.
        parts = self._export()
        name, _chart = self._live_chart(parts)
        self.assertEqual(name, "ppt/charts/chart1.xml")
        self.assertEqual(parts[name], self.source[name])
        self.assertEqual(self._workbook(parts, name), self._workbook(self.source, name))
        self.assertNotIn("native-chart-rebuilt", self.log)

    def test_json_whitespace_change_keeps_exact_source_chart(self) -> None:
        self._import()
        self.metadata.text = json.dumps(self.payload, ensure_ascii=False, indent=2)
        parts = self._export()
        name, _chart = self._live_chart(parts)
        self.assertEqual(name, "ppt/charts/chart1.xml")
        self.assertEqual(parts[name], self.source[name])
        self.assertNotIn("native-chart-rebuilt", self.log)

    def test_point_count_change_rebuilds_with_warning(self) -> None:
        self._import()
        self.payload["categories"].append("Third")
        self.payload["series"][0]["values"].append(56)
        self.payload["series"][0].pop("point_colors", None)
        self.assertIn("native-chart-rebuilt", "\n".join(self._check()))
        parts = self._export()
        name, chart = self._live_chart(parts)
        self.assertNotIn("_data_", name)
        self.assertEqual(chart.find(".//c:val/c:numRef/c:numCache/c:ptCount", NS).get("val"), "3")

    def test_name_edit_adds_missing_series_text_without_rebuilding(self) -> None:
        self._import(named=False)
        self.payload["series"][0]["name"] = "Named now"
        parts = self._export()
        name, chart = self._live_chart(parts)
        self.assertIn("_data_", name)
        self.assertEqual(chart.find(".//c:ser/c:tx/c:v", NS).text, "Named now")
        self.assertIsNotNone(chart.find(".//c:dPt/c:spPr/a:gradFill", NS))

    def test_ancestor_move_with_data_edit_uses_rebuild(self) -> None:
        self._import()
        parent = next(node for node in self.svg.iter() if self.marker in list(node))
        wrapper = ET.Element(f'{{{NS["svg"]}}}g', {"transform": "translate(20 10)"})
        parent.insert(list(parent).index(self.marker), wrapper)
        parent.remove(self.marker)
        wrapper.append(self.marker)
        self.payload["series"][0]["values"] = [21, 43]
        self.assertIn("native-chart-rebuilt", "\n".join(self._check()))

    def test_null_gap_updates_cache_and_blank_workbook_cell(self) -> None:
        self._import()
        self.payload["series"][0]["values"] = [None, 43]
        self.assertEqual(self._check(), [])
        parts = self._export()
        name, chart = self._live_chart(parts)
        cache = chart.find(".//c:val/c:numRef/c:numCache", NS)
        self.assertEqual(cache.find("c:ptCount", NS).get("val"), "2")
        self.assertEqual([point.get("idx") for point in cache.findall("c:pt", NS)], ["1"])
        self.assertIsNone(self._cells(self._workbook(parts, name))["B2"])

    def test_fill_imported_gap_updates_missing_cell(self) -> None:
        self._import(values=(None, 34))
        self.payload["series"][0]["values"] = [21, 34]
        parts = self._export()
        name, chart = self._live_chart(parts)
        self.assertEqual({point.get("idx"): point.find("c:v", NS).text
                          for point in chart.findall(".//c:val/c:numRef/c:numCache/c:pt", NS)},
                         {"0": "21", "1": "34"})
        self.assertEqual(self._cells(self._workbook(parts, name))["B2"], "21")

    def test_shared_string_edit_does_not_change_unrelated_cell(self) -> None:
        self._import()
        # Reuse an existing shared string through a data-only edit, then change
        # the category in a fresh import; the unchanged series name must survive.
        self.payload["categories"][0] = "Revenue"
        parts = self._export()
        chart_part, _ = self._live_chart(parts)
        cells = self._cells(self._workbook(parts, chart_part))
        self.assertEqual(cells["B1"], "Revenue")
        self.assertEqual(cells["A2"], "Revenue")
        # Use the exported PPTX as the next round's immutable source.
        next_workspace = self.root / "next"
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            convert_pptx_to_svg(self.root / "result.pptx", next_workspace,
                                ConvertOptions(inheritance_mode="both", roundtrip=True))
        self.workspace = next_workspace
        self.path = next_workspace / "authoring-svg-flat" / "slide_01.svg"
        self.svg = ET.parse(self.path).getroot()
        self.marker = next(node for node in self.svg.iter() if node.get("data-pptx-replace-with") == "chart")
        self.metadata = self.marker.find('svg:metadata[@type="application/json"]', NS)
        self.payload = json.loads(self.metadata.text)
        self.payload["categories"][0] = "Changed"
        parts = self._export()
        name, _ = self._live_chart(parts)
        cells = self._cells(self._workbook(parts, name))
        self.assertEqual(cells["B1"], "Revenue")
        self.assertEqual(cells["A2"], "Changed")
