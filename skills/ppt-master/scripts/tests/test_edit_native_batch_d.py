#!/usr/bin/env python3
"""Regress Edit Native paint order for new objects on rebuilt round-trip slides.

Usage: python3 -m unittest tests.test_edit_native_batch_d
Dependencies: python-pptx and the converter's normal dependencies.
"""

from __future__ import annotations

import io
import re
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from xml.etree import ElementTree as ET

SCRIPTS_DIR = Path(__file__).resolve().parents[1]
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

from pptx_to_svg.converter import ConvertOptions, convert_pptx_to_svg  # noqa: E402
from svg_to_pptx.pptx_package import cli  # noqa: E402

NS = {
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
}
SOURCE_NAMES = ('base', 'tab', 'callout')
SOURCE_GROUP_IDS = {'base': 'shape-2', 'tab': 'shape-3', 'callout': 'shape-4'}


def _rect(name: str, *, full: bool = False) -> str:
    frame = 'x="0" y="0" width="960" height="720"' if full else 'x="8" y="8" width="100" height="80"'
    return f'<rect id="{name}" {frame} fill="#F12345" />'


class EditNativeBatchDTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def _workspace(self, *, hidden: tuple[str, ...] = ()) -> Path:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for index, name in enumerate(SOURCE_NAMES):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(index + 0.5), Inches(1), Inches(2), Inches(2),
            )
            shape.name = name
            if name in hidden:
                shape._element.find('.//p:cNvPr', NS).set('hidden', '1')
        source = self.root / 'source.pptx'
        prs.save(source)
        project = self.root / 'workspace'
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            convert_pptx_to_svg(source, project, ConvertOptions(inheritance_mode='both', roundtrip=True))
        return project

    def _groups(self, project: Path) -> tuple[Path, str, dict[str, str]]:
        path = project / 'authoring-svg-flat' / 'slide_01.svg'
        text = path.read_text(encoding='utf-8')
        groups = {
            name: re.search(rf'<g id="{group_id}".*?</g>', text, re.S).group(0)
            for name, group_id in SOURCE_GROUP_IDS.items()
            if f'<g id="{group_id}"' in text
        }
        return path, text, groups

    def _author(self, project: Path, *order: str) -> None:
        """Rewrite slide 1 so its top-level objects follow ``order``.

        Source names reuse their imported groups; any other name is a new rect,
        and a ``full:`` prefix makes it full-canvas.
        """
        path, text, groups = self._groups(project)
        first = min(text.index(group) for group in groups.values())
        last = max(text.index(group) + len(group) for group in groups.values())
        body = []
        for name in order:
            if name in groups:
                body.append(groups[name])
            elif name.startswith('full:'):
                body.append(_rect(name.removeprefix('full:'), full=True))
            else:
                body.append(_rect(name))
        path.write_text(text[:first] + '\n'.join(body) + text[last:], encoding='utf-8')

    def _export_roster(self, project: Path) -> tuple[list[str], ET.Element]:
        output = project / 'result.pptx'
        stdout, stderr = io.StringIO(), io.StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            result = cli.main([str(project), '--roundtrip', '-o', str(output)])
        self.assertEqual(result, 0, stdout.getvalue() + stderr.getvalue())
        self.assertIn('rebuilt=1', stdout.getvalue())
        with zipfile.ZipFile(output) as archive:
            slide = ET.fromstring(archive.read('ppt/slides/slide1.xml'))
        tree = slide.find('p:cSld/p:spTree', NS)
        names = [
            child.find('./*/p:cNvPr', NS).get('name')
            for child in tree
            if child.find('./*/p:cNvPr', NS) is not None
        ]
        return names, slide

    def test_replaced_first_object_stays_below_source_objects(self) -> None:
        project = self._workspace()
        self._author(project, 'hr-base-map', 'tab', 'callout')
        names, _slide = self._export_roster(project)
        self.assertEqual(names, ['hr-base-map', 'tab', 'callout'])

    def test_new_objects_keep_leading_middle_and_trailing_positions(self) -> None:
        project = self._workspace()
        self._author(project, 'lead', 'base', 'mid-1', 'mid-2', 'tab', 'callout', 'tail')
        names, _slide = self._export_roster(project)
        self.assertEqual(names, ['lead', 'base', 'mid-1', 'mid-2', 'tab', 'callout', 'tail'])

    def test_reused_slots_anchor_to_the_object_actually_placed(self) -> None:
        project = self._workspace()
        self._author(project, 'callout', 'new', 'tab')
        names, _slide = self._export_roster(project)
        self.assertEqual(names, ['callout', 'new', 'tab'])

    def test_trailing_object_stays_above_unlisted_source_shape(self) -> None:
        project = self._workspace(hidden=('tab',))
        self._author(project, 'base', 'tail')
        names, _slide = self._export_roster(project)
        self.assertEqual(names, ['base', 'tab', 'tail'])

    def test_middle_object_follows_its_authored_predecessor(self) -> None:
        project = self._workspace(hidden=('tab',))
        self._author(project, 'base', 'mid', 'callout')
        names, _slide = self._export_roster(project)
        self.assertEqual(names, ['base', 'mid', 'tab', 'callout'])

    def test_new_full_canvas_rect_stays_a_shape_under_the_source_background(self) -> None:
        project = self._workspace()
        self._author(project, 'full:new-under', 'base', 'tab', 'callout')
        names, slide = self._export_roster(project)
        self.assertEqual(names, ['new-under', 'base', 'tab', 'callout'])
        self.assertIsNone(slide.find('p:cSld/p:bg', NS))
