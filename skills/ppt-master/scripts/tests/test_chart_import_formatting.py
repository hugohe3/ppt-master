#!/usr/bin/env python3
"""PPT Master - Chart Import Formatting Regression Tests

Exercise chart data edits and owning text frames with small native fixtures.

Usage:
    python3 -m unittest tests.test_chart_import_formatting

Examples:
    Run the module from the scripts directory.

Dependencies:
    python-pptx and the local importer, converter, and checker.
"""

from __future__ import annotations

import io
import json
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from pptx import Presentation

from pptx_to_svg.chart_to_svg import extract_native_chart_payload
from pptx_to_svg.color_resolver import ColorPalette
from pptx_to_svg.converter import ConvertOptions, convert_pptx_to_svg
from pptx_to_svg.emu_units import Xfrm
from pptx_to_svg.ooxml_loader import OoxmlPackage, PartRef
from svg_quality.checker import SVGQualityChecker
from svg_to_pptx.drawingml.converter import convert_svg_to_slide_shapes


NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "s": "http://www.w3.org/2000/svg",
}


def _text_style(size: int = 1050, color: str = '<a:schemeClr val="accent1"/>') -> str:
    return (
        '<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr>'
        f'<a:defRPr lang="zh-CN" sz="{size}"><a:solidFill>{color}</a:solidFill>'
        '<a:latin typeface="微软雅黑"/><a:ea typeface="微软雅黑"/>'
        '</a:defRPr></a:pPr><a:endParaRPr lang="zh-CN"/></a:p></c:txPr>'
    )


def _source_chart() -> ET.Element:
    value_style = _text_style(1200, '<a:srgbClr val="123456"/>')
    return ET.fromstring(
        f'<c:chartSpace xmlns:c="{NS["c"]}" xmlns:a="{NS["a"]}">'
        '<c:lang val="zh-CN"/><c:chart><c:plotArea><c:layout><c:manualLayout>'
        '<c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
        '<c:x val="0.1"/><c:y val="0.15"/><c:w val="0.8"/><c:h val="0.65"/>'
        '</c:manualLayout></c:layout><c:barChart><c:barDir val="col"/>'
        '<c:grouping val="clustered"/><c:ser><c:idx val="0"/><c:order val="0"/>'
        '<c:tx><c:v>Budget</c:v></c:tx><c:cat><c:strLit><c:ptCount val="2"/>'
        '<c:pt idx="0"><c:v>Other</c:v></c:pt><c:pt idx="1"><c:v>Equipment</c:v></c:pt>'
        '</c:strLit></c:cat><c:val><c:numLit><c:formatCode>General</c:formatCode>'
        '<c:ptCount val="2"/><c:pt idx="0"><c:v>10</c:v></c:pt>'
        '<c:pt idx="1"><c:v>20</c:v></c:pt></c:numLit></c:val></c:ser>'
        '<c:gapWidth val="40"/><c:overlap val="-27"/><c:axId val="1"/><c:axId val="2"/>'
        '</c:barChart><c:catAx><c:axId val="1"/><c:scaling><c:orientation val="minMax"/>'
        '</c:scaling><c:axPos val="b"/><c:majorTickMark val="in"/><c:tickLblPos val="nextTo"/>'
        f'{_text_style()}<c:crossAx val="2"/><c:crosses val="autoZero"/></c:catAx>'
        '<c:valAx><c:axId val="2"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
        '<c:axPos val="l"/><c:majorGridlines/><c:majorTickMark val="cross"/>'
        '<c:tickLblPos val="nextTo"/>'
        f'{value_style}'
        '<c:crossAx val="1"/><c:crosses val="autoZero"/><c:crossBetween val="between"/>'
        '</c:valAx></c:plotArea><c:legend><c:legendPos val="t"/>'
        f'{_text_style(1350)}</c:legend></c:chart>{_text_style(900)}</c:chartSpace>'
    )


class ChartImportFormattingTests(unittest.TestCase):
    def _import(self, source: ET.Element) -> dict:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "source.pptx"
            with ZipFile(path, "w") as archive:
                archive.writestr("ppt/charts/chart1.xml", ET.tostring(source))
                archive.writestr(
                    "ppt/presentation.xml",
                    '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
                    '<p:sldIdLst/></p:presentation>',
                )
            slide = PartRef("ppt/slides/slide1.xml", ET.Element("slide"), {
                "rId1": {"target": "ppt/charts/chart1.xml"},
            })
            graphic = ET.fromstring(
                f'<a:graphicData xmlns:a="{NS["a"]}" xmlns:c="{NS["c"]}" '
                f'xmlns:r="{NS["r"]}" uri="{NS["c"]}"><c:chart r:id="rId1"/></a:graphicData>'
            )
            palette = ColorPalette(None, None)
            palette.scheme["accent1"] = "347BFE"
            with OoxmlPackage(path) as package:
                result = extract_native_chart_payload(
                    graphic, Xfrm(100, 80, 600, 400), slide, package, palette,
                )
        self.assertIsNone(result.native_status)
        self.assertIsNotNone(result.native_payload)
        return result.native_payload

    def _export(self, payload: dict, *, primary_language: str | None = None) -> ET.Element:
        root = ET.Element(f'{{{NS["s"]}}}svg', {"viewBox": "0 0 1280 720"})
        marker = ET.SubElement(root, f'{{{NS["s"]}}}g', {
            "id": "chart", "data-pptx-replace-with": "chart", "data-pptx-native-authority": "json",
        })
        ET.SubElement(marker, f'{{{NS["s"]}}}metadata', {"type": "application/json"}).text = json.dumps(payload)
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "chart.svg"
            path.write_bytes(ET.tostring(root))
            _, _, _, _, parts, _ = convert_svg_to_slide_shapes(
                path, resource_root=Path(temporary), native_objects=True, primary_language=primary_language,
            )
        charts = [content for name, content in parts.items() if name.startswith("ppt/charts/chart")
                  and name.endswith(".xml") and "/_rels/" not in name]
        self.assertEqual(len(charts), 1)
        return ET.fromstring(charts[0])

    def test_data_edit_preserves_spacing_layout_language_and_axis_typography(self) -> None:
        payload = self._import(_source_chart())
        self.assertEqual(payload.get("gap_width"), 40)
        self.assertEqual(payload.get("overlap"), -27)
        self.assertEqual(payload.get("plot_area"), {"x": 160, "y": 140, "width": 480, "height": 260})
        self.assertEqual(payload.get("lang"), "zh-CN")
        payload["series"][0]["values"] = [60, 140]
        payload["categories"][0] = "Reserve"
        chart = self._export(payload, primary_language="en-US")
        self.assertEqual(chart.find("c:lang", NS).get("val"), "zh-CN")
        self.assertEqual(chart.find(".//c:gapWidth", NS).get("val"), "40")
        self.assertEqual(chart.find(".//c:overlap", NS).get("val"), "-27")
        layout = chart.find(".//c:plotArea/c:layout/c:manualLayout", NS)
        self.assertIsNotNone(layout)
        for key, expected in (("x", 0.1), ("y", 0.15), ("w", 0.8), ("h", 0.65)):
            self.assertAlmostEqual(float(layout.find(f"c:{key}", NS).get("val")), expected)
        for tag, size, color, tick in (("catAx", "1050", "347BFE", "in"),
                                       ("valAx", "1200", "123456", "cross")):
            axis = chart.find(f".//c:{tag}", NS)
            props = axis.find("c:txPr/a:p/a:pPr/a:defRPr", NS)
            self.assertEqual(props.get("sz"), size)
            self.assertEqual(props.find("a:latin", NS).get("typeface"), "Microsoft YaHei")
            self.assertEqual(props.find("a:solidFill/a:srgbClr", NS).get("val"), color)
            self.assertEqual(axis.find("c:majorTickMark", NS).get("val"), tick)
        self.assertEqual(chart.find(".//c:legend/c:txPr//a:defRPr", NS).get("sz"), "1350")

    def test_axis_typography_is_projected_even_when_legacy_axis_gate_passes(self) -> None:
        payload = self._import(_source_chart())
        self.assertEqual(payload.get("axes", {}).get("category", {}).get("font_size"), 14)
        self.assertEqual(payload.get("axes", {}).get("value", {}).get("font_size"), 16)

    def test_identical_point_fills_do_not_enable_vary_colors(self) -> None:
        source = _source_chart()
        plot = source.find(".//c:barChart", NS)
        ET.SubElement(plot, f'{{{NS["c"]}}}varyColors', {"val": "0"})
        series = plot.find("c:ser", NS)
        paint = ET.fromstring(
            f'<c:spPr xmlns:c="{NS["c"]}" xmlns:a="{NS["a"]}">'
            '<a:solidFill><a:srgbClr val="123456"/></a:solidFill></c:spPr>'
        )
        series.append(paint)
        for index in range(2):
            point = ET.SubElement(series, f'{{{NS["c"]}}}dPt')
            ET.SubElement(point, f'{{{NS["c"]}}}idx', {"val": str(index)})
            point.append(ET.fromstring(ET.tostring(paint)))
        payload = self._import(source)
        self.assertNotIn("point_colors", payload["series"][0])
        self.assertEqual(self._export(payload).find(".//c:varyColors", NS).get("val"), "0")

    def test_line_and_marker_colors_keep_distinct_source_paints(self) -> None:
        source = self._export({
            "type": "line", "line_style": "lineMarker", "x": 100, "y": 80, "width": 600, "height": 400,
            "categories": ["A", "B"], "series": [{"name": "Line", "values": [1, 2]}],
            "style": {"colors": ["#112233"]},
        })
        series = source.find(".//c:ser", NS)
        series.find("c:marker/c:spPr/a:solidFill/a:srgbClr", NS).set("val", "FFAA00")
        series.append(ET.fromstring(
            f'<c:dPt xmlns:c="{NS["c"]}" xmlns:a="{NS["a"]}"><c:idx val="0"/>'
            '<c:marker><c:spPr><a:solidFill><a:srgbClr val="00AAFF"/></a:solidFill>'
            '</c:spPr></c:marker></c:dPt>'
        ))
        payload = self._import(source)
        self.assertEqual(payload["style"]["colors"], ["#112233"])
        self.assertEqual(payload["series"][0]["point_colors"], ["#00AAFF", "#FFAA00"])
        rebuilt = self._export(payload).find(".//c:ser", NS)
        self.assertEqual(rebuilt.find("c:spPr/a:ln/a:solidFill/a:srgbClr", NS).get("val"), "112233")
        self.assertEqual([node.get("val") for node in rebuilt.findall(
            "c:dPt/c:marker/c:spPr/a:solidFill/a:srgbClr", NS,
        )], ["00AAFF", "FFAA00"])

    def test_implicit_markers_follow_source_line_color(self) -> None:
        source = self._export({
            "type": "line", "line_style": "lineMarker", "x": 100, "y": 80, "width": 600, "height": 400,
            "categories": ["A", "B"], "series": [{"name": "Line", "values": [1, 2]}],
            "style": {"colors": ["#112233"]},
        })
        marker = source.find(".//c:ser/c:marker", NS)
        marker.remove(marker.find("c:spPr", NS))
        payload = self._import(source)
        self.assertEqual(payload["style"]["colors"], ["#112233"])
        self.assertNotIn("point_colors", payload["series"][0])

    def test_combo_colors_follow_plot_order(self) -> None:
        source = self._export({
            "type": "combo", "x": 100, "y": 80, "width": 600, "height": 400,
            "categories": ["A", "B"], "style": {"colors": ["#112233", "#445566"]},
            "plots": [
                {"type": "column", "series_indices": [1], "series": [{"name": "Bars", "values": [1, 2]}]},
                {"type": "line", "series_indices": [0], "line_style": "lineMarker",
                 "series": [{"name": "Line", "values": [2, 3]}]},
            ],
        })
        payload = self._import(source)
        self.assertEqual(payload["style"]["colors"], ["#112233", "#445566"])
        rebuilt = self._export(payload)
        for plot in ("barChart", "lineChart"):
            path = f".//c:{plot}/c:ser/c:spPr/a:solidFill/a:srgbClr"
            self.assertEqual(rebuilt.find(path, NS).get("val"), source.find(path, NS).get("val"))

    def test_chart_fill_title_and_inherited_axis_style(self) -> None:
        source = _source_chart()
        for axis in source.findall(".//c:catAx", NS) + source.findall(".//c:valAx", NS):
            axis.remove(axis.find("c:txPr", NS))
        for owner, color in ((source, "EEEEEE"), (source.find(".//c:plotArea", NS), "ABCDEF")):
            owner.append(ET.fromstring(
                f'<c:spPr xmlns:c="{NS["c"]}" xmlns:a="{NS["a"]}">'
                f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill></c:spPr>'
            ))
        title = ET.SubElement(source.find("c:chart", NS), f'{{{NS["c"]}}}title')
        title.append(ET.fromstring(
            f'<c:tx xmlns:c="{NS["c"]}" xmlns:a="{NS["a"]}"><c:rich>'
            '<a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr sz="2100">'
            '<a:solidFill><a:srgbClr val="112233"/></a:solidFill>'
            '<a:latin typeface="Arial"/></a:rPr><a:t>Budget</a:t></a:r></a:p></c:rich></c:tx>'
        ))
        chart = self._export(self._import(source))
        self.assertEqual(chart.find("c:spPr/a:solidFill/a:srgbClr", NS).get("val"), "EEEEEE")
        self.assertEqual(chart.find(".//c:plotArea/c:spPr/a:solidFill/a:srgbClr", NS).get("val"), "ABCDEF")
        self.assertEqual(chart.find(".//c:catAx/c:txPr//a:defRPr", NS).get("sz"), "900")
        props = chart.find(".//c:title//a:rPr", NS)
        self.assertEqual(props.get("sz"), "2100")
        self.assertEqual(props.find("a:solidFill/a:srgbClr", NS).get("val"), "112233")

    def test_combo_plot_spacing_and_series_formatting_survive(self) -> None:
        payload = {
            "type": "combo", "x": 100, "y": 80, "width": 600, "height": 400,
            "categories": ["A", "B"], "plots": [
                {"type": "column", "gap_width": 0, "overlap": 0, "series": [{"name": "Bars", "values": [1, 2]}]},
                {"type": "line", "line_style": "lineMarker", "series": [
                    {"name": "Line", "values": [2, 3], "line_width": 4, "marker_size": 12},
                ]},
                {"type": "area", "series": [
                    {"name": "Area", "values": [3, 4], "fill_opacity": 0.25, "line_width": 2},
                ]},
            ],
        }
        imported = self._import(self._export(payload))
        plots = {plot["type"]: plot for plot in imported["plots"]}
        self.assertEqual(plots["column"].get("gap_width"), 0)
        self.assertEqual(plots["column"].get("overlap"), 0)
        self.assertEqual(plots["line"]["series"][0].get("line_width"), 4)
        self.assertEqual(plots["line"]["series"][0].get("marker_size"), 12)
        self.assertEqual(plots["area"]["series"][0].get("fill_opacity"), 0.25)
        self.assertEqual(plots["area"]["series"][0].get("line_width"), 2)

    def test_equal_axis_title_text_keeps_each_axes_own_style(self) -> None:
        source = _source_chart()
        for tag, size, color in (("catAx", "1500", "123456"), ("valAx", "1800", "654321")):
            axis = source.find(f".//c:{tag}", NS)
            axis.append(ET.fromstring(
                f'<c:title xmlns:c="{NS["c"]}" xmlns:a="{NS["a"]}"><c:tx><c:rich>'
                '<a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="900">'
                '<a:latin typeface="Arial"/></a:defRPr></a:pPr>'
                f'<a:r><a:rPr sz="{size}"><a:solidFill><a:srgbClr val="{color}"/>'
                '</a:solidFill></a:rPr><a:t>Shared title</a:t></a:r></a:p></c:rich></c:tx></c:title>'
            ))
        chart = self._export(self._import(source))
        for tag, size, color in (("catAx", "1500", "123456"), ("valAx", "1800", "654321")):
            props = chart.find(f".//c:{tag}/c:title//a:rPr", NS)
            self.assertEqual(props.get("sz"), size)
            self.assertEqual(props.find("a:solidFill/a:srgbClr", NS).get("val"), color)
            self.assertEqual(props.find("a:latin", NS).get("typeface"), "Arial")

    def test_secondary_axis_style_stays_on_secondary_axis(self) -> None:
        payload = {
            "type": "combo", "x": 100, "y": 80, "width": 600, "height": 400,
            "categories": ["A", "B"], "plots": [
                {"type": "column", "series": [{"name": "Bars", "values": [1, 2]}]},
                {"type": "line", "axis": "secondary", "series": [{"name": "Line", "values": [3, 4]}]},
            ],
            "axes": {"secondary_value": {"font_family": "Arial", "font_size": 20, "color": "#123456"}},
        }
        imported = self._import(self._export(payload))
        self.assertEqual(imported["axes"]["secondary_value"]["font_size"], 20)
        self.assertEqual(imported["axes"]["secondary_value"]["color"], "#123456")
        imported["plots"][0]["series"][0]["values"][0] = 10
        chart = self._export(imported)
        axes = chart.findall(".//c:valAx", NS)
        right = next(axis for axis in axes if axis.find("c:axPos", NS).get("val") == "r")
        self.assertEqual(right.find("c:txPr//a:defRPr", NS).get("sz"), "1500")

    def test_point_colors_survive_for_columns_and_line_markers(self) -> None:
        for chart_type, point_colors in (("column", ["#123456", "#ABCDEF"]),
                                          ("line", [None, "#ABCDEF"])):
            with self.subTest(chart_type=chart_type):
                payload = {
                    "type": chart_type, "x": 100, "y": 80, "width": 600, "height": 400,
                    "categories": ["A", "B"], "series": [{
                        "name": "S", "values": [1, 2], "point_colors": point_colors,
                    }],
                }
                if chart_type == "line":
                    payload["series"][0]["marker_size"] = 12
                imported = self._import(self._export(payload))
                self.assertEqual(imported["series"][0].get("point_colors"), point_colors)
                if chart_type == "line":
                    self.assertEqual(imported["series"][0].get("marker_size"), 12)

    def test_scatter_axis_text_overrides_reach_writer(self) -> None:
        payload = {
            "type": "scatter", "x": 100, "y": 80, "width": 600, "height": 400,
            "series": [{"name": "S", "x": [1, 2], "y": [2, 3]}],
            "axes": {"x": {"font_family": "Arial", "font_size": 20,
                           "color": "#123456", "tick_marks": "cross"}},
        }
        chart = self._export(payload)
        props = chart.find(".//c:valAx/c:txPr//a:defRPr", NS)
        self.assertEqual(props.get("sz"), "1500")
        self.assertEqual(chart.find(".//c:valAx/c:majorTickMark", NS).get("val"), "cross")
        imported = self._import(chart)
        self.assertEqual(imported["axes"]["x"].get("font_size"), 20)

    def test_language_default_and_invalid_value(self) -> None:
        payload = {"type": "column", "x": 0, "y": 0, "width": 400, "height": 300,
                   "categories": ["A"], "series": [{"name": "S", "values": [1]}]}
        self.assertEqual(self._export(payload).find("c:lang", NS).get("val"), "en-US")
        self.assertEqual(self._export(payload, primary_language="fr-FR").find("c:lang", NS).get("val"), "fr-FR")
        for invalid in ("", 42, 'en-US"/><injected/>'):
            with self.subTest(invalid=invalid), self.assertRaises(RuntimeError):
                self._export({**payload, "lang": invalid})
        with self.assertRaisesRegex(RuntimeError, "classic charts only"):
            self._export({**payload, "type": "treemap", "values": [1], "levels": [["A"]], "lang": "zh-CN"})


class RoundtripFrameInferenceTests(unittest.TestCase):
    def _check(self, body: str) -> dict:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(source)
            workspace = root / "workspace"
            with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                convert_pptx_to_svg(source, workspace, ConvertOptions(inheritance_mode="both", roundtrip=True))
                page = workspace / "authoring-svg-flat/slide_01.svg"
                svg = ET.parse(page).getroot()
                additions = ET.fromstring(f'<svg xmlns="{NS["s"]}">{body}</svg>')
                svg.extend(additions)
                page.write_bytes(ET.tostring(svg))
                results = SVGQualityChecker().check_roundtrip_workspace(str(workspace))
            self.assertEqual(len(results), 1)
            return results[0]

    def test_group_bounds_prevent_unrelated_root_decoration_frame(self) -> None:
        result = self._check(
            '<rect id="accent" x="10" y="10" width="5" height="50"/>'
            '<g id="caption" data-pptx-bounds="100 100 400 66" transform="translate(20 10)">'
            '<text x="300" y="140" text-anchor="middle" font-family="Arial" font-size="24">Caption</text></g>'
        )
        findings = result["errors"] + result["warnings"]
        self.assertFalse(any("owning frame" in item for item in findings), findings)

    def test_no_frame_does_not_borrow_ancestor_decoration_or_request_metadata(self) -> None:
        result = self._check(
            '<rect id="accent" x="10" y="10" width="5" height="50"/>'
            '<g id="caption"><g><text x="300" y="140" font-family="Arial" font-size="24">Caption</text></g></g>'
        )
        findings = result["errors"] + result["warnings"]
        self.assertFalse(any("exceeds owning frame" in item for item in findings), findings)
        self.assertTrue(any("Cannot verify owning frame" in item for item in findings), findings)
        self.assertFalse(any("add data-pptx-frame" in item or "keep data-pptx-frame" in item for item in findings))

    def test_group_bounds_still_report_real_overflow(self) -> None:
        result = self._check(
            '<g id="caption" data-pptx-bounds="100 100 10 66">'
            '<text x="100" y="140" font-family="Arial" font-size="24">Long caption</text></g>'
        )
        self.assertTrue(any("exceeds owning frame" in item for item in result["errors"]), result)

    def test_direct_rect_frame_requires_text_anchor_inside(self) -> None:
        result = self._check(
            '<g id="caption"><rect id="accent" x="10" y="10" width="5" height="50"/>'
            '<rect id="frame" x="100" y="100" width="400" height="66"/>'
            '<text x="300" y="140" text-anchor="middle" font-family="Arial" font-size="24">Caption</text></g>'
        )
        self.assertFalse(any("owning frame" in item for item in result["errors"] + result["warnings"]), result)

    def test_real_sibling_frame_overflow_stays_advisory(self) -> None:
        result = self._check(
            '<g id="caption"><rect id="frame" x="100" y="100" width="80" height="66"/>'
            '<text x="110" y="140" font-family="Arial" font-size="24">Long caption text</text></g>'
        )
        self.assertFalse(result["errors"], result)
        self.assertTrue(any('exceeds owning frame <rect id="frame">' in item for item in result["warnings"]), result)
        self.assertFalse(any("add data-pptx-frame" in item for item in result["warnings"]), result)


class ChartThemeFontTokenTests(unittest.TestCase):
    def test_theme_tokens_keep_their_own_slots(self) -> None:
        from svg_to_pptx.native_objects.chart_style import _font_face_xml

        self.assertEqual(
            _font_face_xml("+mn-lt, +mn-ea"),
            '<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/><a:cs typeface="+mn-cs"/>',
        )
        self.assertEqual(
            _font_face_xml("+mn-ea"),
            '<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/><a:cs typeface="+mn-cs"/>',
        )
        self.assertIn('<a:ea typeface="+mj-ea"/>', _font_face_xml("Arial, +mj-ea"))
        self.assertIn('<a:latin typeface="Arial"/>', _font_face_xml("Arial, +mj-ea"))

    def test_named_stacks_are_unchanged(self) -> None:
        from svg_to_pptx.native_objects.chart_style import _font_face_xml

        self.assertEqual(
            _font_face_xml("Arial, Microsoft YaHei"),
            '<a:latin typeface="Arial"/><a:ea typeface="Microsoft YaHei"/><a:cs typeface="Arial"/>',
        )
