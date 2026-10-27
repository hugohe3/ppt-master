#!/usr/bin/env python3
"""PPT Master - Hidden Speaker Notes Regression Tests

Keep author-hidden source notes out of editable Markdown and note overlays.

Usage:
    python3 -m unittest tests.test_roundtrip_hidden_notes

Examples:
    Run the module from the scripts directory.

Dependencies:
    python-pptx and the local PPTX import/export pipeline.
"""

from __future__ import annotations

import copy
import io
import tempfile
import unittest
import zipfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation

from pptx_to_svg.converter import ConvertOptions, convert_pptx_to_svg
from svg_to_pptx.pptx_package.cli import main as export_main


NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
NOTES_PART = "ppt/notesSlides/notesSlide1.xml"


class RoundtripHiddenNotesTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "source.pptx"
        self.workspace = self.root / "workspace"

    def _import(self, *, grouped: bool = False, visible: bool = True) -> None:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = "VISIBLE_NOTES"
        deck.save(self.source)
        with zipfile.ZipFile(self.source) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        root = ET.fromstring(parts[NOTES_PART])
        tree = root.find("p:cSld/p:spTree", NS)
        body = next(shape for shape in tree if shape.find(
            'p:nvSpPr/p:nvPr/p:ph[@type="body"]', NS,
        ) is not None)
        hidden = copy.deepcopy(body)
        properties = hidden.find("p:nvSpPr/p:cNvPr", NS)
        properties.set("id", "90")
        properties.set("name", "Hidden note")
        hidden.find(".//a:t", NS).text = "HIDDEN_NOTES"
        if grouped:
            group = ET.fromstring(
                f'<p:grpSp xmlns:p="{NS["p"]}"><p:nvGrpSpPr>'
                '<p:cNvPr id="91" name="Hidden group" hidden="true"/>'
                '<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/></p:grpSp>'
            )
            group.append(hidden)
            tree.append(group)
        else:
            properties.set("hidden", "1")
            tree.append(hidden)
        if not visible:
            tree.remove(body)
        parts[NOTES_PART] = ET.tostring(root)
        with zipfile.ZipFile(self.source, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, payload in parts.items():
                archive.writestr(name, payload)
        self.source_parts = parts
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            self.result = convert_pptx_to_svg(
                self.source, self.workspace, ConvertOptions(inheritance_mode="both", roundtrip=True),
            )
        warnings = [item for item in self.result.diagnostics if item.code == "hidden-notes-content"]
        self.assertEqual(len(warnings), 1)
        self.assertEqual(warnings[0].part_path, NOTES_PART)
        self.assertEqual(warnings[0].shape_id, "90")

    def _export(self) -> dict[str, bytes]:
        output = self.root / "output.pptx"
        stdout, stderr = io.StringIO(), io.StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            result = export_main([str(self.workspace), "--roundtrip", "-o", str(output)])
        self.assertEqual(result, 0, stdout.getvalue() + stderr.getvalue())
        self.assertIn("rebuilt=0", stdout.getvalue())
        with zipfile.ZipFile(output) as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    def test_hidden_shape_does_not_leak_when_visible_notes_are_edited(self) -> None:
        self._import()
        note = self.workspace / "notes/slide_01.md"
        self.assertEqual(note.read_text(encoding="utf-8").strip(), "VISIBLE_NOTES")
        note.write_text(note.read_text(encoding="utf-8") + "\nNOTE_EDIT\n", encoding="utf-8")
        parts = self._export()
        self.assertIn(b"VISIBLE_NOTES", parts[NOTES_PART])
        self.assertIn(b"NOTE_EDIT", parts[NOTES_PART])
        self.assertNotIn(b"HIDDEN_NOTES", parts[NOTES_PART])

    def test_hidden_group_notes_stay_hidden_and_untouched_native_notes_are_exact(self) -> None:
        self._import(grouped=True)
        note = self.workspace / "notes/slide_01.md"
        self.assertNotIn("HIDDEN_NOTES", note.read_text(encoding="utf-8"))
        parts = self._export()
        self.assertEqual(parts[NOTES_PART], self.source_parts[NOTES_PART])
        self.assertEqual(parts["ppt/slides/slide1.xml"], self.source_parts["ppt/slides/slide1.xml"])

    def test_all_hidden_notes_warn_without_creating_markdown(self) -> None:
        self._import(visible=False)
        self.assertFalse((self.workspace / "notes/slide_01.md").exists())
        self.assertEqual(self.result.speaker_notes, ())
        parts = self._export()
        self.assertEqual(parts[NOTES_PART], self.source_parts[NOTES_PART])
