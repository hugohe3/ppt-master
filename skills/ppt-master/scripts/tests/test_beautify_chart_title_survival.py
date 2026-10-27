#!/usr/bin/env python3
"""Keep visible chart titles across PPTX intake and Beautify verification.

Usage:
    python3 -m unittest tests.test_beautify_chart_title_survival

Dependencies:
    python-pptx and the script runtime dependencies.
"""

from __future__ import annotations

import contextlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parents[1]
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import beautify_inventory  # noqa: E402
import pptx_intake  # noqa: E402

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}


class BeautifyChartTitleSurvivalTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "source.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        data = ChartData()
        data.categories = ["CATEGORY_MARKER"]
        data.add_series("SERIES_MARKER", [271])
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1), Inches(1), Inches(8), Inches(5), data,
        ).chart
        chart.has_title = True
        chart.chart_title.text_frame.text = "TITLE_MARKER\n第二行 & <exact>"
        presentation.save(self.source)

    def _intake(self) -> tuple[Path, dict]:
        analysis = self.root / "analysis"
        inventory = analysis / "inventory.json"
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(pptx_intake.main([str(self.source), "-o", str(analysis)]), 0)
            self.assertEqual(beautify_inventory.main([
                str(analysis / "source.slide_library.json"), "-o", str(inventory),
            ]), 0)
        return inventory, json.loads(inventory.read_text(encoding="utf-8"))

    def _mutate_chart(self, mutation) -> None:
        with zipfile.ZipFile(self.source) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        chart = ET.fromstring(entries["ppt/charts/chart1.xml"])
        mutation(chart)
        entries["ppt/charts/chart1.xml"] = ET.tostring(chart, encoding="utf-8")
        with zipfile.ZipFile(self.source, "w") as archive:
            for name, payload in entries.items():
                archive.writestr(name, payload)

    def test_rich_title_survives_intake_and_missing_title_fails_verification(self) -> None:
        inventory, payload = self._intake()
        title = "TITLE_MARKER\n第二行 & <exact>"
        self.assertEqual(payload["slides"][0]["charts"][0]["title"], title)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(beautify_inventory.main([
                str(inventory), "--verify", str(self.source),
            ]), 0)
        presentation = Presentation(self.source)
        presentation.slides[0].shapes[0].chart.has_title = False
        missing_title = self.root / "missing-title.pptx"
        presentation.save(missing_title)
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            self.assertEqual(beautify_inventory.main([
                str(inventory), "--verify", str(missing_title),
            ]), 1)
        self.assertIn("TITLE_MARKER", output.getvalue())

    def test_linked_title_cache_survives_intake(self) -> None:
        def linked_title(root):
            tx = root.find("c:chart/c:title/c:tx", NS)
            tx.clear()
            ref = ET.SubElement(tx, f"{{{NS['c']}}}strRef")
            ET.SubElement(ref, f"{{{NS['c']}}}f").text = "Sheet1!$A$1"
            cache = ET.SubElement(ref, f"{{{NS['c']}}}strCache")
            ET.SubElement(cache, f"{{{NS['c']}}}ptCount", {"val": "1"})
            point = ET.SubElement(cache, f"{{{NS['c']}}}pt", {"idx": "0"})
            ET.SubElement(point, f"{{{NS['c']}}}v").text = "LINKED_TITLE_MARKER"

        self._mutate_chart(linked_title)
        _path, payload = self._intake()
        self.assertEqual(payload["slides"][0]["charts"][0]["title"], "LINKED_TITLE_MARKER")

    def test_title_survives_chart_without_series(self) -> None:
        def remove_series(root):
            for plot in root.findall("c:chart/c:plotArea/*", NS):
                for series in plot.findall("c:ser", NS):
                    plot.remove(series)

        self._mutate_chart(remove_series)
        _path, payload = self._intake()
        chart = payload["slides"][0]["charts"][0]
        self.assertIn("TITLE_MARKER", chart["title"])
        self.assertEqual(chart["series"], [])
